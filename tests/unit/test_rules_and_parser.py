from docx import Document

from core.parser import is_clause_heading, parse_docx
from core.rules import check_penalty_rates, check_terms, normalize_rate_notation, run_rule_engine


def test_rate_notation_normalization_and_penalty_detection():
    assert normalize_rate_notation("每日5‰") == "每日0.5%"
    assert normalize_rate_notation("每日０．３％") == "每日0.3%"
    assert normalize_rate_notation("每日千分之五") == "每日0.5%"
    assert normalize_rate_notation("每日百分之三十") == "每日30%"
    for text in ("逾期每日按5‰支付违约金", "逾期每日按０．３％支付违约金", "逾期每日按千分之五支付违约金"):
        risks = check_penalty_rates(text)
        assert len(risks) == 1, text
        assert risks[0].original_text == text


def test_late_payment_interest_counts_as_breach_remedy():
    samples = [
        "合同金额10000元，乙方应在十日内交付；甲方逾期付款的，应支付逾期利息。",
        "合同金额10000元，乙方应在十日内交付；甲方每日支付逾期付款利息。",
        "合同金额10000元，乙方应在十日内交付；若甲方逾期付款，乙方有权要求甲方支付利息。",
    ]
    for text in samples:
        titles = [r.title for r in check_terms(text)]
        assert "缺少必备条款：违约责任" not in titles


def test_chinese_enumeration_headings_are_parsed_as_clauses(tmp_path):
    assert is_clause_heading("一、合同标的")
    assert is_clause_heading("二. 付款方式")
    path = tmp_path / "chinese-headings.docx"
    doc = Document()
    doc.add_paragraph("测试合同")
    doc.add_paragraph("一、合同标的")
    doc.add_paragraph("提供设备一套。")
    doc.add_paragraph("二、付款方式")
    doc.add_paragraph("验收后付款。")
    doc.save(path)
    parsed = parse_docx(path)
    titles = [c["title"] for c in parsed["clauses"]]
    assert "一、合同标的" in titles
    assert "二、付款方式" in titles
    assert next(c for c in parsed["clauses"] if c["title"] == "一、合同标的")["content"] == "提供设备一套。"


def test_blank_party_names_and_silent_consent_are_deterministic_risks():
    parsed = {
        "full_text": "保密协议\n甲方：\n乙方：\n一方通知变更，对方逾期未答复视为已同意。",
        "clauses": [],
    }
    titles = {risk.title for risk in run_rule_engine(parsed)}
    assert "合同主体名称未填写" in titles
    assert "沉默视为同意的变更或解除机制需明确" in titles
