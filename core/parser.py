"""Word parsing in document order, with explicit source locations and complete batching."""
import re
from typing import Dict, List

CLAUSE_PATTERNS = [r'^第[零〇一二三四五六七八九十百千\d]+[条章]\s*',
                   r'^[一二三四五六七八九十百千]+[、.．]\s*',
                   r'^\d+(?:\.\d+)*(?:[.、\s]|$)', r'^[（(][一二三四五六七八九十\d]+[)）]']


def is_clause_heading(text):
    return any(re.match(p, text.strip()) for p in CLAUSE_PATTERNS)


def parse_docx(file_path) -> Dict:
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc = Document(file_path)
    clauses, full, warnings = [], [], []
    current, pidx, tidx, preceding = None, -1, 0, ''

    def append(title, content, kind='paragraph', indexes=None, rows=None):
        clause = {'clause_id': f'c{len(clauses)+1:03d}', 'title': title,
                  'content': content, 'paragraph_indexes': indexes or [], 'kind': kind,
                  'block_index': len(clauses), 'table_rows': rows or []}
        clauses.append(clause)
        return clause

    for el in doc.element.body.iterchildren():
        if el.tag.endswith('}p'):
            pidx += 1
            p = Paragraph(el, doc)
            text = p.text.strip()
            if not text:
                continue
            full.append(text)
            numbered = p._p.pPr is not None and p._p.pPr.numPr is not None
            if is_clause_heading(text) or numbered or text.startswith('附件'):
                current = append(text, '', indexes=[pidx])
            elif current is None:
                current = append('前言', text, indexes=[pidx])
            else:
                current['content'] = '\n'.join(filter(None, [current['content'], text]))
                current['paragraph_indexes'].append(pidx)
            preceding = text
        elif el.tag.endswith('}tbl'):
            tidx += 1
            table = Table(el, doc)
            rows = []
            for row in table.rows:
                values, previous_cell = [], None
                for cell in row.cells:
                    value = cell.text.strip() if cell._tc is not previous_cell else ''
                    values.append(value)
                    previous_cell = cell._tc
                if any(values):
                    rows.append(values)
            if rows:
                title = f'附件表格：{preceding[:80]}' if preceding.startswith('附件') else f'合同表格{tidx}'
                content = '\n'.join(' | '.join(r) for r in rows)
                append(title, content, 'table', rows=rows)
                full.append(f'【{title}】\n{content}')
                current = None

    extras = set()
    for section in doc.sections:
        for area in (section.header, section.footer):
            text = '\n'.join(p.text.strip() for p in area.paragraphs if p.text.strip())
            if text and text not in extras:
                extras.add(text)
                append('页眉/页脚', text, 'header_footer')
                full.append(text)
    for box in doc.element.xpath('.//w:txbxContent'):
        text = '\n'.join(''.join(p.xpath('.//w:t/text()')) for p in box.xpath('.//w:p'))
        if text.strip():
            append('文本框', text, 'textbox')
            full.append(text)
    if doc.element.xpath('.//w:ins | .//w:del'):
        warnings.append('文档包含未接受的修订，修订层内容可能不在正文中；请提供接受修订后的版本复核。')
    if doc.element.xpath('.//w:drawing | .//w:pict'):
        warnings.append('文档含图片或绘图，本次未识别图片中的文字、印章和签名。')
    if any('footnotes' in r.reltype or 'endnotes' in r.reltype for r in doc.part.rels.values()):
        warnings.append('文档含脚注或尾注，请人工核对附注内容。')
    if not full:
        raise ValueError('文档没有可读取的正文或表格，请提供带文本的 Word 合同。')
    title = next((p.text.strip() for p in doc.paragraphs if p.text.strip()), '未命名合同')
    return {'title': title[:200], 'full_text': '\n'.join(full), 'clauses': clauses, 'warnings': warnings}


def clause_text(clause):
    return f"[{clause['clause_id']}] {clause['title']}\n{clause['content']}"


def clauses_to_text(clauses: List[Dict], max_chars=None) -> str:
    """Never silently drop clauses. Use batch_clauses for model input limits."""
    text = '\n\n'.join(clause_text(c) for c in clauses)
    if max_chars is not None and len(text) > max_chars:
        raise ValueError('条款文本超过指定长度，请使用分批审核，不能截断合同。')
    return text


def batch_clauses(clauses, max_chars=2600):
    """Split oversized clauses while retaining source IDs and every character."""
    if max_chars < 256:
        raise ValueError('分批字符预算过小')
    batches, current, size = [], [], 0
    for c in clauses:
        original = clause_text(c)
        if len(original) <= max_chars:
            pieces = [original]
        else:
            prefix = f"[{c['clause_id']}] （长条款分段）\n"
            width = max_chars - len(prefix)
            source = c['title'] + '\n' + c['content']
            pieces = [prefix + source[i:i+width] for i in range(0, len(source), width)]
        for piece in pieces:
            if current and size + len(piece) + 2 > max_chars:
                batches.append(current)
                current, size = [], 0
            current.append({'clause_id': c['clause_id'], 'text': piece})
            size += len(piece) + 2
    if current:
        batches.append(current)
    return batches

