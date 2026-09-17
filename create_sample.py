"""
生成示例合同：医疗器械委托研发合同
包含多处刻意设计的风险点，用于演示审核系统
"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_sample(output_path: str):
    doc = Document()

    # 标题
    title = doc.add_heading('医疗器械委托研发合同', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 当事人（风险点1：甲方占位符，信息不完整）
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run('甲方（委托方）：').bold = True
    p.add_run('【待填写】医疗科技有限公司')

    p = doc.add_paragraph()
    p.add_run('乙方（受托方）：').bold = True
    p.add_run('某某生物科技有限公司')

    doc.add_paragraph('（以下简称"双方"）')

    # 前言
    doc.add_paragraph(
        '鉴于甲方拟委托乙方进行新型体外诊断试剂的研发工作，双方经友好协商，'
        '根据《中华人民共和国民法典》及相关法律法规，达成如下协议：'
    )

    # 第一条 委托内容
    doc.add_heading('第一条 委托内容', level=1)
    doc.add_paragraph('1.1 甲方委托乙方研发一款用于早期肿瘤筛查的体外诊断试剂产品。')
    doc.add_paragraph('1.2 乙方应按照甲方提出的技术要求完成研发工作。')
    # 风险点：缺少知识产权归属约定

    # 第二条 研发进度
    doc.add_heading('第二条 研发进度', level=1)
    doc.add_paragraph('2.1 乙方应在合同生效后尽快完成研发工作，预计周期约12个月。')
    # 风险点："尽快"模糊，无明确截止日期

    # 第三条 合同金额与支付
    doc.add_heading('第三条 合同金额与支付', level=1)
    doc.add_paragraph('3.1 本合同总金额为￥1,200,000元。')
    # 风险点：只有小写金额，缺少大写
    doc.add_paragraph(
        '3.2 支付方式：合同签订后5个工作日内，甲方支付合同总额的80%作为预付款，即￥960,000元；'
        '剩余20%在产品研发完成后支付。'
    )
    # 风险点：预付款80%过高
    doc.add_paragraph('3.3 乙方收款账户：开户行______，账号______。')
    # 风险点：空白字段

    # 第四条 交付与验收
    doc.add_heading('第四条 交付与验收', level=1)
    doc.add_paragraph('4.1 乙方完成研发后，应向甲方交付研发成果。')
    doc.add_paragraph('4.2 甲方应在收到交付物后及时进行验收，如无异议视为验收合格。')
    # 风险点："及时"无明确期限，验收标准模糊

    # 第五条 双方权利义务
    doc.add_heading('第五条 双方权利义务', level=1)
    doc.add_paragraph('5.1 甲方应按约定支付研发费用。')
    doc.add_paragraph('5.2 乙方应按约定完成研发工作，保证研发过程合法合规。')
    # 风险点：缺少医疗器械资质要求

    # 第六条 保密条款
    doc.add_heading('第六条 保密条款', level=1)
    doc.add_paragraph('6.1 双方应对合作过程中知悉的商业秘密承担保密义务。')
    doc.add_paragraph('6.2 未经对方书面同意，不得向第三方披露保密信息。')
    # 风险点：保密期限不明确

    # 第七条 违约责任
    doc.add_heading('第七条 违约责任', level=1)
    doc.add_paragraph(
        '7.1 如甲方逾期付款，每逾期一日，应按合同总额的5%向乙方支付违约金；'
        '逾期超过30日的，乙方有权解除合同，并要求甲方支付合同总额50%的违约金。'
    )
    # 风险点：违约金过高（日5%、总额50%），且只约束甲方
    doc.add_paragraph('7.2 因乙方原因导致研发失败的，乙方应退还已收取的费用。')

    # 第八条 争议解决
    doc.add_heading('第八条 争议解决', level=1)
    doc.add_paragraph('8.1 因本合同引起的争议，双方应首先通过友好协商解决。')
    doc.add_paragraph(
        '8.2 协商不成的，任何一方均可向北京仲裁委员会申请仲裁，'
        '或向甲方所在地人民法院提起诉讼。'
    )
    # 风险点：或裁或诉，仲裁条款无效

    # 第九条 其他
    doc.add_heading('第九条 其他', level=1)
    doc.add_paragraph('9.1 本合同自双方签字盖章之日起生效。')
    doc.add_paragraph('9.2 本合同一式两份，双方各执一份。')

    # 签署页
    doc.add_paragraph()
    doc.add_paragraph()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = '甲方（盖章）：'
    table.cell(0, 1).text = '乙方（盖章）：'
    table.cell(1, 0).text = '法定代表人（签字）：'
    table.cell(1, 1).text = '法定代表人（签字）：'
    table.cell(2, 0).text = '日期：    年  月  日'
    table.cell(2, 1).text = '日期：    年  月  日'

    doc.save(output_path)

    print("=" * 50)
    print(f"示例合同已生成：{output_path}")
    print("=" * 50)
    print("包含的演示风险点：")
    print("  1. 甲方名称为占位符【待填写】")
    print("  2. 当事人信息不完整（缺少地址、联系方式）")
    print("  3. 合同金额只有小写，缺少大写中文")
    print("  4. 预付款比例高达80%（商业风险）")
    print("  5. 银行账户信息为空白下划线")
    print("  6. 验收期限模糊（'及时'无明确天数）")
    print("  7. 缺少知识产权归属约定（医疗器械重点）")
    print("  8. 缺少医疗器械资质要求")
    print("  9. 保密期限不明确")
    print("  10. 违约金比例过高（日5%、总额50%）")
    print("  11. 违约责任不对等（只约束甲方）")
    print("  12. 争议解决同时约定仲裁和诉讼（无效）")
    print("=" * 50)


if __name__ == "__main__":
    import os
    output = os.path.join(os.path.dirname(__file__), "sample_contract.docx")
    create_sample(output)
